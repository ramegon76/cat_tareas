from pathlib import Path
from collections import OrderedDict
from openpyxl import load_workbook
import json

BASE=Path(__file__).resolve().parents[1]
EXCEL=BASE/'scripts'/'Catálogo_de_tareas.xlsx'
SALIDA=BASE/'data'/'catalogo_tareas.json'
HOJA='Cat_Objetos'

def clean(v):
    return '' if v is None else ' '.join(str(v).replace('\n',' ').replace('\r',' ').split())

wb=load_workbook(EXCEL,read_only=True,data_only=True)
ws=wb[HOJA]
fases=OrderedDict(); vistos=set(); omitidas=[]
for fila,row in enumerate(ws.iter_rows(min_row=4,min_col=5,max_col=12,values_only=True),4):
    fk,fn,ak,an,pk,pn,tk,tn=map(clean,row)
    if not any((fk,fn,ak,an,pk,pn,tk,tn)): continue
    if not all((fk,fn,ak,an,pk,pn,tk,tn)):
        omitidas.append(fila); continue
    f=fases.setdefault(fk,{'clave':fk,'nombre':fn,'actividades':OrderedDict()})
    a=f['actividades'].setdefault(ak,{'clave':ak,'nombre':an,'procedimientos':OrderedDict()})
    p=a['procedimientos'].setdefault(pk,{'clave':pk,'nombre':pn,'tareas':[]})
    ident=(fk,ak,pk,tk)
    if ident not in vistos:
        p['tareas'].append({'clave':tk,'nombre':tn}); vistos.add(ident)

catalogo={'fases':[]}
for f in fases.values():
    fo={'clave':f['clave'],'nombre':f['nombre'],'actividades':[]}
    for a in f['actividades'].values():
        ao={'clave':a['clave'],'nombre':a['nombre'],'procedimientos':[]}
        for p in a['procedimientos'].values():
            ao['procedimientos'].append({'clave':p['clave'],'nombre':p['nombre'],'tareas':p['tareas']})
        fo['actividades'].append(ao)
    catalogo['fases'].append(fo)
SALIDA.write_text(json.dumps(catalogo,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'JSON generado: {SALIDA}')
print(f'Filas incompletas omitidas: {len(omitidas)}')
